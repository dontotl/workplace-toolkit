"""Generate a six-slide public synthetic fixture; never imports customer files."""
from pathlib import Path
import argparse
import io
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def make(output):
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    p.core_properties.author='Workplace Toolkit';p.core_properties.title='Synthetic translation acceptance test'
    def text(shapes,content,x,y,w=11.6,h=1.0,size=26,bold=False):
        box=shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));box.text_frame.word_wrap=True
        for i,line in enumerate(content.split('\n')):
            para=box.text_frame.paragraphs[0] if i==0 else box.text_frame.add_paragraph()
            r=para.add_run();r.text=line;r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string('203F39')
        return box
    titles=['Research to presentation','Keep meaning and emphasis','Evidence in a table','A connected workflow','Translate the whole idea','Ready for review']
    slides=[]
    for i,title in enumerate(titles):
        s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F5F3EA')
        text(s.shapes,title,0.7,0.55,size=34,bold=True)
        badge=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.8), Inches(0.5), Inches(.7), Inches(.7));badge.fill.solid();badge.fill.fore_color.rgb=RGBColor.from_string('C5D8CA');badge.line.fill.background()
        text(s.shapes,str(i+1),11.95,.6,.4,.4,16)
        text(s.shapes,'Synthetic sample — no customer data',.75,6.9,11,.3,12)
        s.notes_slide.notes_text_frame.text='Synthetic presenter note. Keep this note unchanged.'
        slides.append(s)
    text(slides[0].shapes,'One story, supported by evidence.',.8,2,9,1.2,32)
    text(slides[0].shapes,'Collect sources, translate carefully, and explain the next action.',.8,3.5,10,1.6)
    box=text(slides[1].shapes,'',.8,2,11,1.5)
    para=box.text_frame.paragraphs[0]
    for value,bold in [('Preserve ',False),('important terms',True),(' without changing the message.',False)]:
        r=para.add_run();r.text=value;r.font.size=Pt(28);r.font.name='Arial';r.font.bold=bold;r.font.color.rgb=RGBColor.from_string('203F39')
    text(slides[1].shapes,'Oracle AI Database supports the example workflow.',.8,4,11,1.3)
    table=slides[2].shapes.add_table(3,2,Inches(.8),Inches(2),Inches(11.5),Inches(2.5)).table
    for row,values in enumerate([['Measure','Example value'],['Storage','12 GB'],['Response time','250 ms']]):
        for col,value in enumerate(values):
            cell=table.cell(row,col);cell.text=value;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('DDE8DF' if row==0 else 'FFFFFF')
            for para in cell.text_frame.paragraphs:
                for r in para.runs:r.font.size=Pt(24);r.font.name='Arial';r.font.color.rgb=RGBColor.from_string('203F39')
    group=slides[3].shapes.add_group_shape()
    for i,value in enumerate(['Discover sources','Review evidence','Explain the result']):
        card=group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8+i*4.1), Inches(2.5), Inches(3.6), Inches(2))
        card.fill.solid();card.fill.fore_color.rgb=RGBColor.from_string('DDE8DF');card.line.fill.background()
        text(group.shapes,value,1+i*4.1,3,3.2,1,24,True)
    text(slides[4].shapes,'Engineers need implementation detail, while business teams need a clear outcome. A useful presentation connects both perspectives through one consistent story.',.8,2,11.5,3,28)
    text(slides[4].shapes,'Learn more: https://example.com/guide',.8,5.4,11,.8,20)
    image=Image.new('RGB',(320,200),'#DDE8DF');draw=ImageDraw.Draw(image);draw.rounded_rectangle((80,30,240,170),20,fill='#315C49');draw.line((112,95,148,128,214,62),fill='white',width=12)
    stream=io.BytesIO();image.save(stream,format='PNG');stream.seek(0)
    slides[5].shapes.add_picture(stream,Inches(.8),Inches(2),width=Inches(4.5))
    text(slides[5].shapes,'Confirm accuracy before sharing.',6,2,6,1.5,30,True)
    text(slides[5].shapes,'The original file and speaker notes remain unchanged.',6,4,6,1.7,24)
    output.parent.mkdir(parents=True,exist_ok=True);p.save(output)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('output');make(parser.parse_args().output)
