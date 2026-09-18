"""Behavioral regression tests: preserve original package and fail closed."""
import importlib.util
import errno
import subprocess
import stat
import traceback
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches, Pt

SCRIPT = Path(__file__).parents[1] / 'scripts' / 'ppt_bridge.py'


class TranslationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'original.pptx'
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        paragraph = box.text_frame.paragraphs[0]
        r = paragraph.add_run(); r.text = 'Oracle '; r.font.bold = True; r.font.size = Pt(24)
        r = paragraph.add_run(); r.text = 'supports 12 teams.'; r.font.size = Pt(20)
        s.notes_slide.notes_text_frame.text = 'Private note stays local.'
        p.save(self.source)

    def bridge(self):
        self.assertTrue(SCRIPT.exists(), 'extract/apply/verify bridge is missing')
        spec = importlib.util.spec_from_file_location('ppt_bridge', SCRIPT)
        module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        return module

    def data(self, b):
        m = b.extract(self.source)
        t = {'source_sha256': m['source_sha256'], 'manifest_sha256':m.get('manifest_sha256'), 'translations': [
            {'id': m['segments'][0]['id'], 'text': 'Oracle '},
            {'id': m['segments'][1]['id'], 'text': '12개 팀을 지원합니다.'}]}
        return m, t

    def test_extract_excludes_notes_and_has_context(self):
        b = self.bridge(); m, _ = self.data(b)
        self.assertEqual(m['slide_count'], 1)
        self.assertEqual(len(m['segments']), 2)
        self.assertEqual(m['segments'][0]['context'], 'Oracle supports 12 teams.')

    def test_apply_preserves_notes_media_runs_and_original(self):
        b = self.bridge(); m,t = self.data(b); original = self.source.read_bytes()
        out = self.root / 'translated.pptx'; b.apply(self.source, m, t, out)
        self.assertEqual(self.source.read_bytes(), original)
        with zipfile.ZipFile(self.source) as z, zipfile.ZipFile(out) as new:
            for name in z.namelist():
                if name != 'ppt/slides/slide1.xml': self.assertEqual(z.read(name), new.read(name), name)
        p = Presentation(out); runs=p.slides[0].shapes[0].text_frame.paragraphs[0].runs
        self.assertTrue(runs[0].font.bold); self.assertEqual(runs[0].font.size.pt,24)
        self.assertEqual(runs[1].font.size.pt,20); self.assertIn('12개',runs[1].text)
        self.assertTrue(b.verify(self.source,out,m,t)['passed'])

    def test_missing_duplicate_unknown_translation_rejected(self):
        b=self.bridge();m,t=self.data(b)
        for entries in [t['translations'][:1],t['translations']+[t['translations'][0]],t['translations']+[{'id':'fake','text':'x'}]]:
            with self.subTest(entries=entries), self.assertRaises(ValueError):
                b.apply(self.source,m,dict(t,translations=entries),self.root/'out.pptx')

    def test_source_hash_and_manifest_tampering_rejected(self):
        b=self.bridge();m,t=self.data(b)
        m['segments'][0]['source']='Changed'
        with self.assertRaises(ValueError): b.apply(self.source,m,t,self.root/'out.pptx')
        m,_=self.data(b);t['source_sha256']='bad'
        with self.assertRaises(ValueError): b.apply(self.source,m,t,self.root/'out.pptx')

    def test_number_url_and_product_changes_rejected(self):
        b=self.bridge();m,t=self.data(b)
        for text in ['13개 팀을 지원합니다.', '팀을 지원합니다.']:
            t['translations'][1]['text']=text
            with self.assertRaises(ValueError): b.apply(self.source,m,t,self.root/'out.pptx')
        m,t=self.data(b);m=b.extract(self.source,['Oracle']);t['manifest_sha256']=m['manifest_sha256'];t['translations'][0]['text']='다른 제품 '
        with self.assertRaises(ValueError):b.apply(self.source,m,t,self.root/'out.pptx')

    def test_existing_output_and_source_never_overwritten(self):
        b=self.bridge();m,t=self.data(b);out=self.root/'out.pptx';out.write_bytes(b'existing')
        with self.assertRaises(FileExistsError):b.apply(self.source,m,t,out)
        self.assertEqual(out.read_bytes(),b'existing')
        with self.assertRaises((ValueError,FileExistsError)):b.apply(self.source,m,t,self.source)

    def test_failed_save_does_not_publish_partial(self):
        b=self.bridge();m,t=self.data(b);out=self.root/'out.pptx'
        with patch.object(b.os,'link',side_effect=OSError('disk full')):
            with self.assertRaises(OSError):b.apply(self.source,m,t,out)
        self.assertFalse(out.exists())
        self.assertEqual(list(self.root.glob('.ppt-bridge-*')),[])

    def test_verify_detects_changed_notes(self):
        b=self.bridge();m,t=self.data(b);out=self.root/'out.pptx';b.apply(self.source,m,t,out)
        broken=self.root/'broken.pptx'
        with zipfile.ZipFile(out) as src,zipfile.ZipFile(broken,'w') as dest:
            for n in src.namelist():dest.writestr(n,src.read(n)+b' ' if n=='ppt/notesSlides/notesSlide1.xml' else src.read(n))
        self.assertFalse(b.verify(self.source,broken,m,t)['passed'])

    def test_protected_tokens_keep_sign_currency_order_and_allow_korean_particles(self):
        b=self.bridge()
        for before,after in [('-12%','+12%'),('$12','€12'),('10 to 20','20에서 10'),('https://example.com/a','https://example.com/b')]:
            with self.subTest(before=before):self.assertNotEqual(b.protected(before,[]),b.protected(after,[]))
        self.assertEqual(b.protected('12 GB and 250 ms',[]),b.protected('12 GB를 사용하고 250 ms의 지연',[]))

    def test_manifest_warning_and_policy_tampering_rejected(self):
        b=self.bridge();m,t=self.data(b)
        m['warnings']=[]
        m['protected_terms']=['Oracle']
        with self.assertRaises(ValueError):b.apply(self.source,m,t,self.root/'out.pptx')
        m,t=self.data(b);m['warnings']=[{'slide':1,'kind':'made-up'}]
        with self.assertRaises(ValueError):b.apply(self.source,m,t,self.root/'out.pptx')

    def test_master_layout_gaps_are_reported(self):
        b=self.bridge();m=b.extract(self.source)
        self.assertTrue(any(w['kind']=='master-layout-text-not-translated' for w in m['warnings']))

    def test_unsupported_hardlinks_use_safe_exclusive_copy(self):
        b=self.bridge();m,t=self.data(b);out=self.root/'portable.pptx'
        with patch.object(b.os,'link',side_effect=OSError(errno.ENOTSUP,'unsupported')):
            b.apply(self.source,m,t,out)
        self.assertTrue(b.verify(self.source,out,m,t)['passed'])

    def test_legacy_conversion_error_is_content_free(self):
        b=self.bridge();source=self.root/'customer-secret.ppt';source.write_bytes(b'legacy')
        with patch.object(b.shutil,'which',return_value='soffice'),patch.object(b.subprocess,'run',side_effect=subprocess.CalledProcessError(1,['soffice',str(source)])):
            with self.assertRaises(RuntimeError) as error:b.convert_legacy(source,self.root/'out.pptx')
        self.assertNotIn('customer-secret',str(error.exception))

    def test_legacy_traceback_does_not_expose_command(self):
        b=self.bridge();source=self.root/'customer-secret.ppt';source.write_bytes(b'legacy')
        with patch.object(b.shutil,'which',return_value='soffice'),patch.object(b.subprocess,'run',side_effect=subprocess.CalledProcessError(1,['soffice',str(source)])):
            try:b.convert_legacy(source,self.root/'out.pptx')
            except RuntimeError as exc:
                self.assertIsNone(exc.__cause__)
                self.assertNotIn('customer-secret',traceback.format_exc())
            else:self.fail('conversion failure expected')

    def test_url_cannot_move_out_of_hyperlinked_run(self):
        p=Presentation(self.source);runs=p.slides[0].shapes[0].text_frame.paragraphs[0].runs
        runs[0].text='https://example.com';runs[0].hyperlink.address='https://example.com';runs[1].text=' More'
        p.save(self.source);b=self.bridge();m=b.extract(self.source)
        t={'source_sha256':m['source_sha256'],'manifest_sha256':m['manifest_sha256'],'translations':[
            {'id':m['segments'][0]['id'],'text':'자세히: '},{'id':m['segments'][1]['id'],'text':'https://example.com'}]}
        with self.assertRaises(ValueError):b.apply(self.source,m,t,self.root/'out.pptx')

    def test_fallback_keeps_private_permissions(self):
        if __import__('os').name=='nt':self.skipTest('POSIX permission bits')
        b=self.bridge();m,t=self.data(b);out=self.root/'private.pptx'
        with patch.object(b.os,'link',side_effect=OSError(errno.ENOTSUP,'unsupported')):b.apply(self.source,m,t,out)
        self.assertEqual(stat.S_IMODE(out.stat().st_mode),0o600)

    def test_unchanged_media_is_streamed(self):
        with zipfile.ZipFile(self.source,'a') as z:z.writestr('ppt/media/video1.mp4',b'local test media')
        b=self.bridge();m,t=self.data(b);read=zipfile.ZipFile.read
        def guarded_read(archive,name,*args,**kwargs):
            filename=getattr(name,'filename',name)
            if filename.startswith('ppt/media/'):raise AssertionError('media must be streamed')
            return read(archive,name,*args,**kwargs)
        with patch.object(zipfile.ZipFile,'read',guarded_read):
            b.apply(self.source,m,t,self.root/'out.pptx')

if __name__=='__main__':unittest.main()
